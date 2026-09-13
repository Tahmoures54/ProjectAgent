# Path: pms_app/utils/excel_io.py
"""Shared helpers for Excel import / export (openpyxl)."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from pms_app.utils.jalali import parse_jalali_to_gregorian


def require_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation

        return {
            "Workbook": Workbook,
            "load_workbook": load_workbook,
            "Alignment": Alignment,
            "Font": Font,
            "PatternFill": PatternFill,
            "Border": Border,
            "Side": Side,
            "get_column_letter": get_column_letter,
            "DataValidation": DataValidation,
        }
    except ImportError:
        return None


def workbook_to_bytes(wb) -> bytes:
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def xlsx_response(content: bytes, filename: str):
    from flask import Response

    return Response(
        content,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    for ch in ("\u200c", " ", "_", "-", ".", ":", "،", "\n", "\t"):
        text = text.replace(ch, "")
    return text


def header_index(headers: Sequence[Any], aliases: Dict[str, Sequence[str]]) -> Dict[str, int]:
    """Map logical keys to column indexes using Persian/English aliases."""
    normalized = [normalize_header(h) for h in headers]
    alias_map: Dict[str, str] = {}
    for key, names in aliases.items():
        for name in names:
            alias_map[normalize_header(name)] = key
    idx: Dict[str, int] = {}
    for i, header in enumerate(normalized):
        key = alias_map.get(header)
        if key and key not in idx:
            idx[key] = i
    return idx


def cell(row: Sequence[Any], idx: Dict[str, int], key: str, default=None):
    i = idx.get(key)
    if i is None or i >= len(row):
        return default
    val = row[i]
    if isinstance(val, str):
        val = val.strip()
        return val if val != "" else default
    return default if val is None else val


def parse_excel_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            # Excel serial date (Windows epoch)
            serial = float(value)
            if 1 <= serial < 100000:
                return (datetime(1899, 12, 30) + timedelta(days=serial)).date()
        except (OverflowError, ValueError):
            pass
    text = str(value).strip()
    if not text:
        return None
    iso = text.replace("/", "-")[:10]
    if len(iso) == 10 and iso[4] == "-" and iso[7] == "-" and iso[:4].isdigit():
        year = int(iso[:4])
        if year >= 1700:
            try:
                return date.fromisoformat(iso)
            except ValueError:
                pass
    jalali = parse_jalali_to_gregorian(text)
    if jalali is not None:
        return jalali
    try:
        if "T" in text or " " in text:
            return datetime.fromisoformat(text.replace("Z", "")[:19]).date()
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def to_int(value: Any, default: Optional[int] = None) -> Optional[int]:
    if value is None or value == "":
        return default
    try:
        return int(float(str(value).replace(",", "").replace("٬", "").translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789"))))
    except (TypeError, ValueError):
        return default


def to_float(value: Any, default: Optional[float] = None) -> Optional[float]:
    if value is None or value == "":
        return default
    try:
        return float(str(value).replace(",", "").replace("٬", "").translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")))
    except (TypeError, ValueError):
        return default


def to_decimal(value: Any) -> Optional[Decimal]:
    f = to_float(value)
    if f is None:
        return None
    try:
        return Decimal(str(f))
    except (InvalidOperation, ValueError):
        return None


def style_header_row(ws, row: int = 1, fill_hex: str = "0E7F9B"):
    ox = require_openpyxl()
    if not ox:
        return
    fill = ox["PatternFill"]("solid", fgColor=fill_hex)
    font = ox["Font"](bold=True, color="FFFFFF")
    align = ox["Alignment"](horizontal="center", vertical="center", wrap_text=True)
    thin = ox["Border"](
        left=ox["Side"](style="thin", color="D1D5DB"),
        right=ox["Side"](style="thin", color="D1D5DB"),
        top=ox["Side"](style="thin", color="D1D5DB"),
        bottom=ox["Side"](style="thin", color="D1D5DB"),
    )
    for cell_obj in ws[row]:
        cell_obj.fill = fill
        cell_obj.font = font
        cell_obj.alignment = align
        cell_obj.border = thin
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = f"A{row + 1}"
    ws.sheet_view.rightToLeft = True


def autosize(ws, min_width: int = 12, max_width: int = 42):
    ox = require_openpyxl()
    if not ox:
        return
    for col in ws.columns:
        letter = ox["get_column_letter"](col[0].column)
        width = min_width
        for c in col[:40]:
            if c.value is not None:
                width = max(width, min(max_width, len(str(c.value)) + 2))
        ws.column_dimensions[letter].width = width


def iter_data_rows(ws) -> Tuple[List[Any], Iterable[Tuple[int, Tuple[Any, ...]]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = list(rows[0])
    data = []
    for i, row in enumerate(rows[1:], start=2):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        data.append((i, tuple(row)))
    return headers, data
