# tests/test_daily_reports.py
from __future__ import annotations

from datetime import date, timedelta

from pms_app.models import Company, Contract, ContractItem, DailyReport, Project, Role, User
from pms_app.models.project_membership import ProjectMembership


def _role(session, name):
    return Role.query.filter_by(name=name).first()


def _user(session, prefix, role_name, company_id=None, password="pass1234"):
    user = User(
        email=f"{prefix}@example.com",
        full_name=prefix,
        is_active=True,
        company_id=company_id,
    )
    user.set_password(password)
    role = _role(session, role_name)
    if role:
        user.roles = [role]
    session.add(user)
    session.commit()
    return user


def _company_project(session, suffix="DR"):
    company = Company(name=f"Co {suffix}")
    session.add(company)
    session.flush()
    project = Project(
        company_id=company.id,
        project_code=f"PRJ-{suffix}",
        project_name=f"Project {suffix}",
        industry="construction",
        base_currency="IRR",
        status="active",
        finish_date=date.today() - timedelta(days=3),
    )
    session.add(project)
    session.commit()
    return company, project


def _login(client, email, password="pass1234"):
    return client.post("/login", data={"email": email, "password": password})


def test_daily_reports_list_requires_login(client):
    response = client.get("/daily-reports/")
    assert response.status_code in (302, 401)


def test_daily_report_model_workflow_and_progress(db_session):
    company, project = _company_project(db_session, "WF")
    admin = _user(db_session, "dr_admin", "company_admin", company.id)
    worker = _user(db_session, "dr_worker", "contractor", company.id)

    contract = Contract(
        company_id=company.id,
        project_id=project.id,
        contract_number="CNT-DR",
        contract_title="Main",
        contract_type="EPC",
        pricing_model="lumpsum",
        currency="IRR",
        status="active",
    )
    db_session.add(contract)
    db_session.flush()
    item = ContractItem(
        company_id=company.id,
        contract_id=contract.id,
        title="Foundation",
        status="open",
        actual_progress_percentage=10,
        original_amount=1000,
        adjusted_amount=1000,
    )
    db_session.add(item)
    db_session.commit()

    report = DailyReport(
        company_id=company.id,
        project_id=project.id,
        report_date=date.today(),
        submitted_by_id=worker.id,
        work_performed="بتن‌ریزی فونداسیون",
        progress_updates=[{"contract_item_id": item.id, "progress_percent": 40}],
        status="draft",
    )
    db_session.add(report)
    db_session.commit()

    report.submit(worker.id)
    assert report.status == "submitted"
    report.approve(admin.id, comment="تأیید", apply_progress=True)
    db_session.commit()

    assert report.status == "approved"
    assert report.progress_applied is True
    assert float(item.actual_progress_percentage) == 40.0


def test_contractor_submits_and_admin_approves_via_http(client, db_session):
    company, project = _company_project(db_session, "HTTP")
    admin = _user(db_session, "http_admin", "company_admin", company.id)
    worker = _user(db_session, "http_worker", "contractor", company.id)
    db_session.add(
        ProjectMembership(
            project_id=project.id,
            user_id=worker.id,
            role="contractor",
            status="active",
        )
    )
    db_session.commit()

    _login(client, worker.email)
    create = client.post(
        f"/daily-reports/project/{project.id}/new",
        data={
            "report_date": date.today().isoformat(),
            "work_performed": "نصب اسکلت",
            "manpower_total": "8",
            "action": "submit",
        },
        follow_redirects=False,
    )
    assert create.status_code in (301, 302)
    report = DailyReport.query.filter_by(project_id=project.id).first()
    assert report is not None
    assert report.status == "submitted"

    client.get("/logout")
    _login(client, admin.email)
    review = client.post(
        f"/daily-reports/{report.id}/review",
        data={"action": "approve", "comment": "ok", "apply_progress": "yes"},
        follow_redirects=False,
    )
    assert review.status_code in (301, 302)
    db_session.refresh(report)
    assert report.status == "approved"


def test_daily_report_excel_template_and_import(client, db_session):
    from io import BytesIO

    from openpyxl import Workbook, load_workbook

    company, project = _company_project(db_session, "XLS")
    admin = _user(db_session, "xls_admin", "company_admin", company.id)
    db_session.add(
        ProjectMembership(
            project_id=project.id,
            user_id=admin.id,
            role="admin",
            status="active",
        )
    )
    db_session.commit()
    _login(client, admin.email)

    template = client.get(f"/daily-reports/project/{project.id}/template.xlsx")
    assert template.status_code == 200
    assert "spreadsheetml" in template.content_type
    wb = load_workbook(BytesIO(template.data))
    assert "گزارش روزانه" in wb.sheetnames
    assert "مهندسی" in wb.sheetnames

    out = Workbook()
    ws = out.active
    ws.title = "گزارش روزانه"
    ws.append(["تاریخ", "شرح کار", "نیروی انسانی", "فاز EPC", "هوا"])
    ws.append([date.today().isoformat(), "لوله کشی واحد ۳", 15, "اجرا", "آفتابی"])
    mp = out.create_sheet("نیروی انسانی")
    mp.append(["تاریخ", "نقش", "تعداد"])
    mp.append([date.today().isoformat(), "جوشکار", 5])
    eng = out.create_sheet("مهندسی")
    eng.append(["تاریخ", "شماره مدرک", "عنوان", "وضعیت"])
    eng.append([date.today().isoformat(), "PID-003", "P&ID واحد ۳", "IFC"])
    bio = BytesIO()
    out.save(bio)
    bio.seek(0)

    uploaded = client.post(
        f"/daily-reports/project/{project.id}/import",
        data={"file": (bio, "daily.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=False,
    )
    assert uploaded.status_code in (301, 302)
    report = DailyReport.query.filter_by(project_id=project.id).first()
    assert report is not None
    assert report.status == "draft"
    assert report.work_performed == "لوله کشی واحد ۳"
    assert report.manpower_total == 15
    assert report.epc_phase == "construction"
    assert report.manpower_details and report.manpower_details[0]["role"] == "جوشکار"
    assert report.engineering_outputs and report.engineering_outputs[0]["doc_no"] == "PID-003"

    exported = client.get(f"/daily-reports/project/{project.id}/export.xlsx")
    assert exported.status_code == 200
    assert "spreadsheetml" in exported.content_type


def test_epc_controls_page(client, db_session):
    company, project = _company_project(db_session, "EPC")
    admin = _user(db_session, "epc_admin", "company_admin", company.id)
    db_session.add(
        ProjectMembership(
            project_id=project.id,
            user_id=admin.id,
            role="admin",
            status="active",
        )
    )
    db_session.commit()
    _login(client, admin.email)
    page = client.get(f"/projects/{project.id}/epc")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "اتاق کنترل" in html
    assert "مهندسی" in html
